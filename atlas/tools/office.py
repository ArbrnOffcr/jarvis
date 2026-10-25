"""Word, Excel und PowerPoint erstellen und bearbeiten.

Neue Dokumente landen im Ausgabeordner (Standard: DATENARCHIV/JARVIS-Ausgaben)
und werden nur gemeldet. Bestehende Dateien zu ändern braucht eine Bestätigung;
vorher wird eine Sicherungskopie angelegt."""
import re
import shutil
import subprocess
from datetime import datetime
from pathlib import Path

from atlas.tools import S, Tool, ToolError

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


class Office:
    def __init__(self, cfg: dict, archive: Path | None, shared: Path):
        self.root = archive
        self.out = (archive / cfg["office"]["ausgabe"]) if archive else None
        self.backup = shared / "_Sicherungen"

    # ---------- Pfade ----------
    def _target(self, pfad: str, endung: str) -> Path:
        """Neue Datei: nur Dateiname -> Ausgabeordner. Unterordner sind erlaubt,
        aber nur innerhalb des Ausgabeordners."""
        if not self.out:
            raise ToolError("DATENARCHIV nicht gefunden")
        name = Path(pfad.strip().lstrip("/")).with_suffix(endung)
        p = (self.out / name).resolve() if not str(name).startswith(self.out.name + "/") \
            else (self.root / name).resolve()
        if self.out.resolve() not in p.parents:
            raise ToolError("Neue Dokumente nur im Ausgabeordner – danach ggf. mit datei_verschieben ablegen")
        if p.exists():
            stamp = datetime.now().strftime("%Y-%m-%d_%H%M")
            p = p.with_name(f"{p.stem}_{stamp}{p.suffix}")
        p.parent.mkdir(parents=True, exist_ok=True)
        return p

    def _existing(self, pfad: str, endungen: tuple) -> Path:
        if not self.root:
            raise ToolError("DATENARCHIV nicht gefunden")
        rel = pfad.strip().lstrip("/").removeprefix("DATENARCHIV/")
        p = (self.root / rel).resolve()
        if self.root.resolve() not in p.parents or not p.exists():
            raise ToolError(f"'{pfad}' nicht gefunden")
        if p.suffix.lower() not in endungen:
            raise ToolError(f"Erwartet {', '.join(endungen)}")
        return p

    def _sichern(self, p: Path) -> str:
        rel = p.relative_to(self.root.resolve())
        ziel = self.backup / datetime.now().strftime("%Y-%m-%d_%H%M%S_%f") / rel
        ziel.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(p, ziel)
        return str(ziel.relative_to(self.backup.parent))

    def _rel(self, p: Path) -> str:
        return str(p.relative_to(self.root.resolve()))

    # ---------- Word ----------
    @staticmethod
    def _runs(par, text: str):
        for part in INLINE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                par.add_run(part[2:-2]).bold = True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                par.add_run(part[1:-1]).italic = True
            else:
                par.add_run(part)

    def _markdown(self, doc, md: str):
        """Einfaches Markdown -> Word: Überschriften, Listen, Tabellen, fett/kursiv."""
        lines = md.replace("\r", "").split("\n")
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()
            if line.startswith("|"):                       # Tabelle
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                        rows.append(cells)
                    i += 1
                if rows:
                    t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                    t.style = "Table Grid"
                    for r, row in enumerate(rows):
                        for cidx, val in enumerate(row):
                            cell = t.cell(r, cidx)
                            cell.text = ""
                            self._runs(cell.paragraphs[0], val)
                            if r == 0:
                                for run in cell.paragraphs[0].runs:
                                    run.bold = True
                continue
            m = re.match(r"^(#{1,4})\s+(.*)", line)
            if m:
                doc.add_heading(m.group(2), level=len(m.group(1)))
            elif re.match(r"^\s*[-*•]\s+", line):
                self._runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*•]\s+", "", line))
            elif re.match(r"^\s*\d+[.)]\s+", line):
                self._runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line))
            elif line.strip():
                self._runs(doc.add_paragraph(), line)
            i += 1

    def word_erstellen(self, dateiname: str, titel: str, inhalt: str) -> dict:
        import docx
        from docx.shared import Pt
        p = self._target(dateiname, ".docx")
        doc = docx.Document()
        st = doc.styles["Normal"]
        st.font.name, st.font.size = "Calibri", Pt(11)
        if titel:
            doc.add_heading(titel, level=0)
        self._markdown(doc, inhalt)
        doc.save(p)
        return {"erstellt": self._rel(p), "art": "word"}

    def word_ergaenzen(self, pfad: str, inhalt: str) -> dict:
        import docx
        p = self._existing(pfad, (".docx",))
        sicherung = self._sichern(p)
        doc = docx.Document(p)
        self._markdown(doc, inhalt)
        doc.save(p)
        return {"ergaenzt": self._rel(p), "sicherung": sicherung}

    def word_ersetzen(self, pfad: str, suchen: str, ersetzen: str) -> dict:
        import docx
        p = self._existing(pfad, (".docx",))
        doc = docx.Document(p)
        pars = list(doc.paragraphs) + [par for t in doc.tables for row in t.rows
                                       for c in row.cells for par in c.paragraphs]
        n = 0
        for par in pars:
            if suchen in par.text:
                # Formatierung des ersten Abschnitts beibehalten, Text zusammenführen
                text = par.text.replace(suchen, ersetzen)
                n += par.text.count(suchen)
                for r in par.runs[1:]:
                    r.text = ""
                if par.runs:
                    par.runs[0].text = text
        if not n:
            raise ToolError(f"'{suchen}' kommt im Dokument nicht vor")
        sicherung = self._sichern(p)
        doc.save(p)
        return {"ersetzt": n, "datei": self._rel(p), "sicherung": sicherung}

    # ---------- Excel ----------
    @staticmethod
    def _fill(ws, kopf, zeilen):
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        if kopf:
            ws.append(kopf)
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F4E79")
            ws.freeze_panes = "A2"
        for z in zeilen or []:
            ws.append(z)
        for idx, col in enumerate(ws.columns, 1):
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(idx)].width = min(60, max(10, width + 2))

    def excel_erstellen(self, dateiname: str, blaetter: list) -> dict:
        import openpyxl
        p = self._target(dateiname, ".xlsx")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for b in blaetter or [{"name": "Tabelle1"}]:
            ws = wb.create_sheet(str(b.get("name", "Tabelle"))[:31])
            self._fill(ws, b.get("kopf"), b.get("zeilen"))
        wb.save(p)
        return {"erstellt": self._rel(p), "art": "excel", "blaetter": len(wb.sheetnames)}

    def excel_zellen(self, pfad: str, zellen: dict, blatt: str = "") -> dict:
        import openpyxl
        p = self._existing(pfad, (".xlsx", ".xlsm"))
        sicherung = self._sichern(p)
        wb = openpyxl.load_workbook(p)
        ws = wb[blatt] if blatt else wb.active
        for ref, val in (zellen or {}).items():
            ws[ref] = val
        wb.save(p)
        return {"geaendert": len(zellen or {}), "datei": self._rel(p), "blatt": ws.title,
                "sicherung": sicherung}

    # ---------- PowerPoint ----------
    def pptx_erstellen(self, dateiname: str, titel: str, folien: list, untertitel: str = "") -> dict:
        import pptx
        from pptx.util import Pt
        p = self._target(dateiname, ".pptx")
        prs = pptx.Presentation()
        prs.slide_width, prs.slide_height = pptx.util.Inches(13.333), pptx.util.Inches(7.5)
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = titel
        if len(s.placeholders) > 1:
            s.placeholders[1].text = untertitel or datetime.now().strftime("%d.%m.%Y")
        for f in folien or []:
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = f.get("titel", "")
            body = s.placeholders[1].text_frame
            body.clear()
            for k, punkt in enumerate(f.get("punkte", [])):
                par = body.paragraphs[0] if k == 0 else body.add_paragraph()
                ebene = len(punkt) - len(punkt.lstrip("-")) if punkt.startswith("-") else 0
                par.text, par.level = punkt.lstrip("- ").strip(), min(ebene, 3)
                for r in par.runs:
                    r.font.size = Pt(22 if ebene == 0 else 18)
            if f.get("notizen"):
                s.notes_slide.notes_text_frame.text = f["notizen"]
        prs.save(p)
        return {"erstellt": self._rel(p), "art": "powerpoint", "folien": len(prs.slides)}

    # ---------- öffnen ----------
    def oeffnen(self, pfad: str) -> dict:
        p = self._existing(pfad, (".docx", ".xlsx", ".xlsm", ".pptx", ".pdf", ".md", ".txt", ".csv"))
        try:
            subprocess.run(["open", str(p)], check=True, timeout=15)
        except (OSError, subprocess.SubprocessError) as e:
            raise ToolError(f"Konnte nicht öffnen: {e}")
        return {"geoeffnet": self._rel(p)}


def tools(cfg, secrets, archive, shared) -> list[Tool]:
    o = Office(cfg, archive, shared)
    blatt = {"type": "object", "properties": {
        "name": S("Blattname"),
        "kopf": {"type": "array", "items": {"type": "string"}, "description": "Spaltenüberschriften"},
        "zeilen": {"type": "array", "items": {"type": "array", "items": {}},
                   "description": "Zeilen; Zahlen als Zahl, Formeln als Text mit '=' (z. B. '=SUM(B2:B9)')"}}}
    folie = {"type": "object", "properties": {
        "titel": S("Folientitel"),
        "punkte": {"type": "array", "items": {"type": "string"},
                   "description": "Stichpunkte; Unterpunkt mit führendem '-'"},
        "notizen": S("Sprechernotizen, optional")}}
    return [
        Tool("word_erstellen", "Neues Word-Dokument (Brief, Antrag, Bericht) aus Markdown erstellen "
             "(#-Überschriften, Listen, |Tabellen|, **fett**). Ablage im Ausgabeordner.",
             "office.anlegen", o.word_erstellen,
             {"dateiname": S("Dateiname ohne Pfad, z. B. 'Antrag_Elternzeit'"), "titel": S("Dokumenttitel"),
              "inhalt": S("Inhalt als Markdown")}, ["dateiname", "inhalt"],
             label=lambda a: f"Word-Dokument „{a['dateiname']}“ erstellen"),
        Tool("word_ergaenzen", "Text (Markdown) an ein bestehendes Word-Dokument anhängen.",
             "office.aendern", o.word_ergaenzen, {"pfad": S("Pfad im DATENARCHIV"), "inhalt": S("Markdown")},
             ["pfad", "inhalt"], label=lambda a: f"„{a['pfad']}“ ergänzen (mit Sicherungskopie)"),
        Tool("word_ersetzen", "In einem Word-Dokument Text suchen und ersetzen.", "office.aendern",
             o.word_ersetzen, {"pfad": S("Pfad"), "suchen": S("Text"), "ersetzen": S("Neuer Text")},
             ["pfad", "suchen", "ersetzen"], label=lambda a: f"In „{a['pfad']}“ „{a['suchen']}“ ersetzen"),
        Tool("excel_erstellen", "Neue Excel-Arbeitsmappe mit einem oder mehreren Blättern erstellen "
             "(Kopfzeile formatiert, Formeln möglich).", "office.anlegen", o.excel_erstellen,
             {"dateiname": S("Dateiname ohne Pfad"), "blaetter": {"type": "array", "items": blatt}},
             ["dateiname", "blaetter"], label=lambda a: f"Excel-Datei „{a['dateiname']}“ erstellen"),
        Tool("excel_zellen", "Zellen in einer bestehenden Excel-Datei setzen (z. B. {\"B4\": 120, \"C10\": \"=SUM(C2:C9)\"}).",
             "office.aendern", o.excel_zellen,
             {"pfad": S("Pfad"), "zellen": {"type": "object", "description": "Zelle -> Wert"},
              "blatt": S("Blattname, optional")}, ["pfad", "zellen"],
             label=lambda a: f"{len(a.get('zellen', {}))} Zellen in „{a['pfad']}“ ändern"),
        Tool("pptx_erstellen", "Neue PowerPoint-Präsentation (16:9) mit Titelfolie und Stichpunktfolien.",
             "office.anlegen", o.pptx_erstellen,
             {"dateiname": S("Dateiname ohne Pfad"), "titel": S("Titel"), "untertitel": S("optional"),
              "folien": {"type": "array", "items": folie}}, ["dateiname", "titel", "folien"],
             label=lambda a: f"Präsentation „{a['dateiname']}“ erstellen"),
        Tool("datei_oeffnen", "Datei auf dem Mac in der passenden App öffnen (Word, Excel, Keynote …).",
             "datei.oeffnen", o.oeffnen, {"pfad": S("Pfad im DATENARCHIV")}, ["pfad"],
             label=lambda a: f"„{a['pfad']}“ öffnen"),
    ]
