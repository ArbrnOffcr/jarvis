"""Gesprächsaufnahme mit Transkript.
Das Dashboard schickt das Mikrofon (16 kHz, 16 Bit, mono) über /aufnahme; der Kern schreibt eine
WAV-Datei (nur lokal, ~/.atlas/aufnahmen), transkribiert sie danach mit Whisper auf dem Mac und legt
das Transkript als Word-Datei in DATENARCHIV/<Ausgabeordner>/Transkripte ab. Nichts geht in die Cloud."""
import threading
import time
import wave
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np

from atlas.paths import LOCAL_DIR

RATE = 16000
DIR = LOCAL_DIR / "aufnahmen"
MAX_STUNDEN = 4
AUFBEWAHREN_TAGE = 30


def _zeit(s: float) -> str:
    s = int(s)
    return f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}" if s >= 3600 else f"{s // 60:02d}:{s % 60:02d}"


class Aufnahme:
    def __init__(self):
        DIR.mkdir(parents=True, exist_ok=True)
        self._aufraeumen()
        self.start = datetime.now()
        self.pfad = DIR / f"Aufnahme_{self.start:%Y-%m-%d_%H-%M-%S}.wav"
        self.wav = wave.open(str(self.pfad), "wb")
        self.wav.setnchannels(1)
        self.wav.setsampwidth(2)
        self.wav.setframerate(RATE)
        self.samples = 0
        self.offen = True

    @property
    def sekunden(self) -> float:
        return self.samples / RATE

    def schreiben(self, data: bytes) -> bool:
        """False, wenn die Höchstdauer erreicht ist."""
        if not self.offen:
            return False
        self.wav.writeframes(data)
        self.samples += len(data) // 2
        return self.sekunden < MAX_STUNDEN * 3600

    def schliessen(self) -> None:
        if self.offen:
            self.offen = False
            self.wav.close()

    @staticmethod
    def _aufraeumen():
        grenze = time.time() - AUFBEWAHREN_TAGE * 86400
        for f in DIR.glob("Aufnahme_*.wav"):
            try:
                if f.stat().st_mtime < grenze:
                    f.unlink()
            except OSError:
                pass


def _absaetze(segmente: list[dict]) -> list[tuple[float, str]]:
    """Whisper-Segmente zu lesbaren Absätzen: neuer Absatz bei Pause > 1,5 s oder nach ~45 s."""
    out, start, teile, ende = [], None, [], 0.0
    for s in segmente:
        text = (s.get("text") or "").strip()
        if not text:
            continue
        if start is not None and (s["start"] - ende > 1.5 or s["start"] - start > 45):
            out.append((start, " ".join(teile)))
            start, teile = None, []
        if start is None:
            start = s["start"]
        teile.append(text)
        ende = s["end"]
    if teile:
        out.append((start, " ".join(teile)))
    # Whisper wiederholt bei Stille manchmal denselben Satz – Dubletten hintereinander entfernen
    sauber = []
    for t, text in out:
        if not sauber or sauber[-1][1] != text:
            sauber.append((t, text))
    return sauber


def transkribieren(wav_pfad: Path, whisper_modell: str, lock: threading.Lock) -> dict:
    with wave.open(str(wav_pfad), "rb") as w:
        pcm = np.frombuffer(w.readframes(w.getnframes()), dtype=np.int16)
    audio = pcm.astype(np.float32) / 32768.0
    import mlx_whisper
    with lock:
        r = mlx_whisper.transcribe(audio, path_or_hf_repo=whisper_modell,
                                   condition_on_previous_text=False, verbose=None)
    return {"sprache": r.get("language") or "de", "absaetze": _absaetze(r.get("segments") or []),
            "dauer": len(pcm) / RATE}


def auswerten(ergebnis: dict, lokal) -> dict:
    """Zusammenfassung, Entscheidungen und Aufgaben – mit dem lokalen Modell, nichts verlässt den Mac.
    Lange Gespräche werden abschnittsweise gelesen und danach zusammengeführt."""
    from atlas import hilfen
    text = "\n".join(t for _, t in ergebnis.get("absaetze") or [])
    leer = {"zusammenfassung": [], "entscheidungen": [], "aufgaben": []}
    if len(text) < 200 or not lokal or not lokal.available():
        return leer
    heute = datetime.now().strftime("%d.%m.%Y")
    stuecke = [text[i:i + 9000] for i in range(0, len(text), 9000)][:8]
    gesamt = {"zusammenfassung": [], "entscheidungen": [], "aufgaben": []}
    for i, st in enumerate(stuecke, 1):
        j = hilfen.lokal_json(
            lokal, "Du wertest Gesprächstranskripte aus. Antworte nur mit JSON, erfinde nichts.",
            f"Heute ist der {heute}. Abschnitt {i}/{len(stuecke)} eines Gesprächs. Antworte nur mit JSON: "
            '{"zusammenfassung": ["2-5 kurze Stichpunkte"], "entscheidungen": ["…"], '
            '"aufgaben": [{"titel": "konkrete Aufgabe, mit Verb", "wer": "Person oder leer", '
            '"faellig": "JJJJ-MM-TT oder leer"}]}. Nur Aufgaben, die wirklich vereinbart wurden.\n\n' + st,
            900) or {}
        for k in ("zusammenfassung", "entscheidungen"):
            gesamt[k] += [str(x).strip()[:200] for x in (j.get(k) or []) if str(x).strip()]
        for a in j.get("aufgaben") or []:
            if isinstance(a, dict) and str(a.get("titel", "")).strip():
                gesamt["aufgaben"].append({"titel": str(a["titel"]).strip()[:140],
                                           "wer": str(a.get("wer") or "").strip()[:40],
                                           "faellig": str(a.get("faellig") or "").strip()[:10]})
    # Dubletten entfernen
    gesehen, aufgaben = set(), []
    for a in gesamt["aufgaben"]:
        k = a["titel"].lower()[:50]
        if k not in gesehen:
            gesehen.add(k); aufgaben.append(a)
    gesamt["aufgaben"] = aufgaben[:15]
    gesamt["zusammenfassung"] = gesamt["zusammenfassung"][:10]
    gesamt["entscheidungen"] = gesamt["entscheidungen"][:8]
    return gesamt


def als_word(ergebnis: dict, start: datetime, ziel_ordner: Path, titel: str = "",
             auswertung: dict | None = None) -> Path:
    from docx import Document
    from docx.shared import Pt, RGBColor
    ziel_ordner.mkdir(parents=True, exist_ok=True)
    name = f"{'Telefonnotiz' if 'telefon' in (titel or '').lower() else 'Transkript'}_{start:%Y-%m-%d_%H-%M}"
    ziel = ziel_ordner / f"{name}.docx"
    n = 2
    while ziel.exists():
        ziel = ziel_ordner / f"{name}_{n}.docx"
        n += 1
    doc = Document()
    doc.styles["Normal"].font.name = "Helvetica Neue"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(titel or "Gesprächstranskript", level=1)
    ende = start + timedelta(seconds=ergebnis["dauer"])
    meta = doc.add_table(rows=0, cols=2)
    for k, v in (("Datum", f"{start:%d.%m.%Y}"), ("Uhrzeit", f"{start:%H:%M} – {ende:%H:%M} Uhr"),
                 ("Dauer", _zeit(ergebnis["dauer"])), ("Sprache", ergebnis["sprache"]),
                 ("Erstellt", "automatisch durch JARVIS (Whisper, lokal auf dem Mac)")):
        row = meta.add_row().cells
        row[0].text, row[1].text = k, v
    doc.add_paragraph()
    aw = auswertung or {}
    for kopf, key in (("Zusammenfassung", "zusammenfassung"), ("Entscheidungen", "entscheidungen")):
        if aw.get(key):
            doc.add_heading(kopf, level=2)
            for x in aw[key]:
                doc.add_paragraph(x, style="List Bullet")
    if aw.get("aufgaben"):
        doc.add_heading("Aufgaben", level=2)
        for a in aw["aufgaben"]:
            zusatz = ", ".join(x for x in (a.get("wer"), a.get("faellig")) if x)
            doc.add_paragraph(a["titel"] + (f" ({zusatz})" if zusatz else ""), style="List Bullet")
    if any(aw.get(k) for k in ("zusammenfassung", "entscheidungen", "aufgaben")):
        doc.add_heading("Wortlaut", level=2)
    if not ergebnis["absaetze"]:
        doc.add_paragraph("(Keine Sprache erkannt.)")
    for t, text in ergebnis["absaetze"]:
        p = doc.add_paragraph()
        z = p.add_run(f"[{_zeit(t)}]  ")
        z.font.size = Pt(9)
        z.font.color.rgb = RGBColor(0x6B, 0x7C, 0x8C)
        p.add_run(text)
    hinweis = doc.add_paragraph()
    r = hinweis.add_run("Hinweis: Maschinell erstellt – Namen, Zahlen und Fachbegriffe vor Weitergabe prüfen.")
    r.italic = True
    r.font.size = Pt(9)
    doc.save(ziel)
    return ziel
