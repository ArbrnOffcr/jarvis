"""Steuer-Sammelmappe: Belege übers Jahr vormerken, im Frühjahr fertige Liste für den Steuerberater.

Belege kommen automatisch aus dem Posteingang (lokal erkannt) oder per Chat („merk das für die
Steuer vor“). Die Dateien bleiben, wo sie sind – die Mappe ist eine Liste mit Verweisen. Beim Export
entsteht eine Excel-Liste im Ausgabeordner, auf Wunsch mit Kopien der Belege nach Kategorien."""
import re
import shutil
import threading
import uuid
from datetime import date, datetime
from pathlib import Path

from atlas import hilfen
from atlas.paths import read_private_json, write_private_json
from atlas.tools import S, Tool, ToolError

ZUSTAND = "steuer.json"

# Kategorie → Schlüsselwörter (Wortanfang, kleingeschrieben). Ein starkes Wort genügt,
# sonst braucht es zwei Treffer.
KATEGORIEN = {
    "Vermietung (Anlage V)": (["mieter", "mietvertrag", "mieteinnahme", "nebenkostenabrechnung", "betriebskosten",
                              "grundsteuer", "hausverwaltung", "mietobjekt", "vermiet"],
                             ["nebenkostenabrechnung", "grundsteuer", "vermiet"]),
    "Handwerker & haushaltsnahe Dienste (§ 35a)": (["handwerk", "reparatur", "wartung", "montage", "schornsteinfeger",
                                                   "heizungs", "sanitär", "elektro", "maler", "gärtner",
                                                   "reinigung", "arbeitslohn", "lohnanteil", "anfahrt"],
                                                  ["schornsteinfeger", "lohnanteil", "arbeitslohn"]),
    "Kinderbetreuung": (["kita", "kindertagesstätte", "kindergarten", "hort", "tagesmutter", "tagespflege",
                         "betreuungsbeitrag", "elternbeitrag", "betreuungskosten"],
                        ["betreuungsbeitrag", "elternbeitrag", "tagesmutter", "kindertagesstätte"]),
    "Vorsorge & Versicherungen": (["beitragsbescheinigung", "jahresbescheinigung", "versicherungsbeitrag",
                                   "haftpflicht", "unfallversicherung", "berufsunfähig", "lebensversicherung",
                                   "pflegeversicherung", "krankenversicherung", "rentenversicherung"],
                                  ["beitragsbescheinigung", "jahresbescheinigung"]),
    "Werbungskosten": (["fortbildung", "fachliteratur", "arbeitsmittel", "bundeswehrverband", "trennungsgeld",
                        "reisekosten", "dienstreise", "umzugskosten", "doppelte haushaltsführung", "seminar",
                        "lehrgang", "berufsverband", "gewerkschaft"],
                       ["trennungsgeld", "umzugskosten", "doppelte haushaltsführung", "berufsverband"]),
    "Krankheitskosten": (["eigenanteil", "zuzahlung", "brille", "zahnersatz", "heilmittel", "apotheke"],
                         ["zahnersatz"]),
    "Spenden & Beiträge": (["zuwendungsbestätigung", "spendenquittung", "spendenbescheinigung", "spende",
                            "mitgliedsbeitrag"],
                           ["zuwendungsbestätigung", "spendenquittung", "spendenbescheinigung"]),
    "Finanzamt & Bescheide": (["finanzamt", "einkommensteuer", "steuerbescheid", "vorauszahlung", "steuernummer",
                               "festsetzung"],
                              ["steuerbescheid", "einkommensteuerbescheid", "vorauszahlung"]),
}


def erkennen(text: str) -> str | None:
    t = (text or "").lower()
    best, punkte = None, 0
    for kat, (worte, stark) in KATEGORIEN.items():
        p = sum(1 for w in worte if re.search(r"\b" + re.escape(w), t))
        if any(re.search(r"\b" + re.escape(w), t) for w in stark):
            p += 2
        if p > punkte:
            best, punkte = kat, p
    return best if punkte >= 2 else None


def _kat(kategorie: str) -> str:
    k = (kategorie or "").lower()
    for name in KATEGORIEN:
        if k and (k in name.lower() or name.lower().startswith(k)):
            return name
    raise ToolError("Kategorie unbekannt. Möglich: " + "; ".join(KATEGORIEN))


def _steuerjahr() -> int:
    """Bis Ende Juli gilt meist noch das Vorjahr als „aktuelle“ Erklärung."""
    heute = date.today()
    return heute.year - 1 if heute.month <= 7 else heute.year


class Steuer:
    def __init__(self, core):
        self.c = core
        self.lock = threading.Lock()

    @property
    def archiv(self) -> Path | None:
        return self.c.archiv.archive if self.c.archiv else None

    def zustand(self) -> dict:
        z = read_private_json(ZUSTAND, {})
        z.setdefault("belege", [])
        return z

    def vormerken(self, pfad: str, kategorie: str, jahr: int = 0, betrag: str = "", titel: str = "",
                  quelle: str = "chat") -> dict:
        kat = _kat(kategorie)
        a = self.archiv
        if not a or not (a / pfad.strip("/")).exists():
            raise ToolError(f"Datei „{pfad}“ nicht im DATENARCHIV gefunden")
        pfad = pfad.strip("/")
        if not jahr:
            m = re.match(r"^(\d{4})[_-]", Path(pfad).name)
            jahr = int(m.group(1)) if m else date.today().year
        b = None
        if betrag:
            from atlas.workflows.beihilfe_parser import betrag as lese_betrag
            b = lese_betrag(str(betrag) if "," in str(betrag) else str(betrag).replace(".", ","))
        with self.lock:
            z = self.zustand()
            alt = next((x for x in z["belege"] if x["pfad"] == pfad and x["kategorie"] == kat), None)
            if alt:
                return {"vorhanden": True, **alt}
            e = {"id": uuid.uuid4().hex[:8], "jahr": int(jahr), "kategorie": kat, "pfad": pfad,
                 "titel": (titel or Path(pfad).stem)[:120], "betrag": b, "quelle": quelle,
                 "erfasst": datetime.now().isoformat(timespec="minutes")}
            z["belege"].append(e)
            write_private_json(ZUSTAND, z)
        self.c.audit.log("steuer", aktion="vorgemerkt", jahr=jahr, kategorie=kat)
        return e

    def entfernen(self, beleg: str) -> dict:
        with self.lock:
            z = self.zustand()
            vorher = len(z["belege"])
            z["belege"] = [x for x in z["belege"] if x["id"] != beleg and x["pfad"] != beleg.strip("/")]
            if len(z["belege"]) == vorher:
                raise ToolError("Beleg nicht gefunden")
            write_private_json(ZUSTAND, z)
        return {"entfernt": beleg}

    def mappe(self, jahr: int = 0) -> dict:
        jahr = int(jahr or _steuerjahr())
        gruppen: dict[str, dict] = {}
        for x in self.zustand()["belege"]:
            if x["jahr"] != jahr:
                continue
            g = gruppen.setdefault(x["kategorie"], {"anzahl": 0, "summe": 0.0, "belege": []})
            g["anzahl"] += 1
            g["summe"] = round(g["summe"] + (x["betrag"] or 0), 2)
            g["belege"].append({"id": x["id"], "titel": x["titel"], "pfad": x["pfad"], "betrag": x["betrag"]})
        fehlt = [k for k in KATEGORIEN if k not in gruppen and k != "Finanzamt & Bescheide"]
        return {"jahr": jahr, "kategorien": gruppen, "belege_gesamt": sum(g["anzahl"] for g in gruppen.values()),
                "noch_ohne_beleg": fehlt}

    def export(self, jahr: int = 0, belege_kopieren: bool = False) -> dict:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        jahr = int(jahr or _steuerjahr())
        a = self.archiv
        if not a:
            raise ToolError("DATENARCHIV nicht gefunden")
        m = self.mappe(jahr)
        if not m["belege_gesamt"]:
            raise ToolError(f"Für {jahr} sind noch keine Belege vorgemerkt")
        ordner = a / self.c.cfg["office"]["ausgabe"] / "Steuer"
        ordner.mkdir(parents=True, exist_ok=True)
        ziel = hilfen.frei(ordner / f"Steuer-Belege {jahr}.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Belege {jahr}"
        kopf = ["Kategorie", "Beleg", "Betrag €", "Datei (DATENARCHIV)"]
        ws.append(kopf)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E79")
        for kat, g in sorted(m["kategorien"].items()):
            for b in g["belege"]:
                ws.append([kat, b["titel"], b["betrag"], b["pfad"]])
            ws.append([f"Summe {kat}", "", g["summe"] or None, ""])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
        for col, w in zip("ABCD", (38, 48, 12, 80)):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
            for c in row:
                c.number_format = '#,##0.00 "€"'
        wb.save(ziel)
        kopiert = 0
        if belege_kopieren:
            basis = ordner / f"{jahr} Belege"
            for kat, g in m["kategorien"].items():
                for b in g["belege"]:
                    q = a / b["pfad"]
                    if q.exists():
                        zd = basis / hilfen.dateiname(kat, 60)
                        zd.mkdir(parents=True, exist_ok=True)
                        shutil.copy2(q, hilfen.frei(zd / q.name))
                        kopiert += 1
        rel = str(ziel.relative_to(a))
        return {"erstellt": rel, "art": "xlsx", "belege": m["belege_gesamt"], "kopiert": kopiert}

    def briefing_hinweis(self) -> str:
        """Im Februar/März einmal wöchentlich: Mappe ist bereit."""
        heute = date.today()
        if heute.month not in (2, 3) or heute.weekday() != 0:
            return ""
        m = self.mappe(heute.year - 1)
        if not m["belege_gesamt"]:
            return ""
        return (f"Steuer {heute.year - 1}: {m['belege_gesamt']} Belege in der Sammelmappe – "
                f"„Exportiere die Steuer-Sammelmappe“ erstellt die Liste für den Steuerberater.")


def tools(core) -> list[Tool]:
    s = core.steuer
    return [
        Tool("steuer_vormerken",
             "Eine Datei im DATENARCHIV als Steuerbeleg in die Sammelmappe aufnehmen (Datei bleibt, wo sie ist).",
             "steuer.schreiben", s.vormerken,
             {"pfad": S("Pfad relativ zum DATENARCHIV"), "kategorie": S("Kategorie", enum=list(KATEGORIEN)),
              "jahr": {"type": "integer", "description": "Steuerjahr, optional"},
              "betrag": S("Betrag, optional, z. B. '120,50'"), "titel": S("Kurzbezeichnung, optional")},
             ["pfad", "kategorie"], label=lambda a: f"Steuerbeleg vormerken ({a.get('kategorie')})"),
        Tool("steuer_mappe", "Inhalt der Steuer-Sammelmappe eines Jahres: Belege je Kategorie mit Summen.",
             "steuer.lesen", s.mappe, {"jahr": {"type": "integer", "description": "Steuerjahr, optional"}}),
        Tool("steuer_export",
             "Steuer-Sammelmappe als Excel-Liste für den Steuerberater erstellen (Ausgabeordner), "
             "optional mit Kopien aller Belege nach Kategorien.",
             "office.anlegen", s.export,
             {"jahr": {"type": "integer", "description": "Steuerjahr, optional"},
              "belege_kopieren": {"type": "boolean", "description": "Belege mitkopieren"}},
             label=lambda a: f"Steuer-Sammelmappe {a.get('jahr') or ''} exportieren".replace("  ", " ")),
        Tool("steuer_entfernen", "Beleg aus der Steuer-Sammelmappe nehmen (id oder Pfad).",
             "steuer.schreiben", s.entfernen, {"beleg": S("id oder Pfad")}, ["beleg"]),
    ]
