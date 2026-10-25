"""Text aus Dateien holen – inkl. Texterkennung (Apple Vision) für Scans und Fotos."""
import subprocess
from pathlib import Path

TEXT = {".txt", ".md", ".csv", ".tsv", ".json", ".log"}
PDF = {".pdf"}
WORD = {".docx"}
EXCEL = {".xlsx", ".xlsm"}
PPT = {".pptx"}
TEXTUTIL = {".doc", ".rtf", ".odt", ".rtfd"}
BILD = {".jpg", ".jpeg", ".png", ".heic", ".tif", ".tiff"}
SUPPORTED = TEXT | PDF | WORD | EXCEL | PPT | TEXTUTIL | BILD

MAX_PDF_PAGES = 300


class ExtractError(Exception):
    pass


def extract(path: Path, ocr: bool = True) -> list[tuple[int, str]]:
    """Liefert [(seite, text), ...]; Seite 0 = ohne Seitenbezug."""
    ext = path.suffix.lower()
    if ext in TEXT:
        return [(0, _read_text(path))]
    if ext in PDF:
        return _pdf(path, ocr)
    if ext in WORD:
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return [(0, "\n".join(parts))]
    if ext in EXCEL:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        out = []
        for i, ws in enumerate(wb.worksheets[:5], 1):
            rows = []
            for r, row in enumerate(ws.iter_rows(values_only=True)):
                if r >= 2000:
                    break
                vals = [str(v) for v in row if v not in (None, "")]
                if vals:
                    rows.append(" | ".join(vals))
            out.append((i, f"[Blatt {ws.title}]\n" + "\n".join(rows)))
        wb.close()
        return out
    if ext in PPT:
        import pptx
        prs = pptx.Presentation(str(path))
        out = []
        for i, slide in enumerate(prs.slides, 1):
            texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
            out.append((i, "\n".join(texts)))
        return out
    if ext in TEXTUTIL:
        r = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(path)],
                           capture_output=True, text=True, timeout=120)
        if r.returncode:
            raise ExtractError(r.stderr[-200:])
        return [(0, r.stdout)]
    if ext in BILD:
        if not ocr:
            return []
        from PIL import Image
        if ext == ".heic":
            from pillow_heif import register_heif_opener
            register_heif_opener()
        with Image.open(path) as img:
            return [(0, _ocr(img.convert("RGB")))]
    raise ExtractError(f"Dateityp {ext} nicht unterstützt")


def _read_text(path: Path) -> str:
    raw = path.read_bytes()[:5_000_000]
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _pdf(path: Path, ocr: bool) -> list[tuple[int, str]]:
    import fitz  # PyMuPDF
    out = []
    with fitz.open(str(path)) as doc:
        for i, page in enumerate(doc, 1):
            if i > MAX_PDF_PAGES:
                break
            text = page.get_text().strip()
            if len(text) < 30 and ocr:          # vermutlich Scan
                from PIL import Image
                pix = page.get_pixmap(dpi=200)
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                text = _ocr(img)
            out.append((i, text))
    return out


def _ocr(img) -> str:
    """Texterkennung mit Apple Vision (offline, Deutsch + Englisch)."""
    try:
        from ocrmac import ocrmac
    except ImportError:
        return ""
    res = ocrmac.OCR(img, language_preference=["de-DE", "en-US"],
                     recognition_level="accurate").recognize()
    # Ergebnisse: (text, konfidenz, box) – Box y von unten; nach Zeilen sortieren
    res.sort(key=lambda r: (-round(r[2][1], 2), r[2][0]))
    return "\n".join(r[0] for r in res)


def chunk(pages: list[tuple[int, str]], size: int = 1200, overlap: int = 150) -> list[dict]:
    """Text in überlappende Abschnitte teilen (für die Suche)."""
    out = []
    for seite, text in pages:
        text = " ".join(text.split())
        if not text:
            continue
        start = 0
        while start < len(text):
            end = min(len(text), start + size)
            if end < len(text):                 # an Wortgrenze schneiden
                cut = text.rfind(" ", start + size // 2, end)
                end = cut if cut > 0 else end
            out.append({"seite": seite, "text": text[start:end]})
            if end >= len(text):
                break
            start = max(end - overlap, start + 1)
    return out
